"""Renders a resume JSON (matching the master.json schema) to a clean, ATS-friendly single-
column PDF: Jinja2 -> HTML -> Playwright page.pdf(). No tables-for-layout, no images, real
selectable text throughout (per the spec's explicit no-other-PDF-library rule).

Also renders a matching .docx via python-docx, purely so it can be opened and edited in Apple
Pages (File > Open, or double-click) - there's no library that writes native .pages files
(it's a proprietary bundle format only Pages itself can produce reliably), so .docx is the
practical bridge: Pages opens it natively and "Save As > Pages" is one click from there."""
from __future__ import annotations

import base64
import functools
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Pt, RGBColor
from jinja2 import Environment, FileSystemLoader
from playwright.sync_api import sync_playwright

from .. import config

TEMPLATE_DIR = config.RESUMES_DIR / "templates"
TEMPLATE_NAME = "resume.html.j2"
FONT_DIR = TEMPLATE_DIR / "fonts"


@functools.lru_cache(maxsize=1)
def _font_b64() -> dict[str, str]:
    """Fonts are embedded as base64 data URIs (not linked by relative path) because
    page.set_content() has no document base URL to resolve relative paths against."""
    names = {"font_regular": "Poppins-Regular.ttf", "font_medium": "Poppins-Medium.ttf",
            "font_semibold": "Poppins-SemiBold.ttf", "font_bold": "Poppins-Bold.ttf"}
    return {key: base64.b64encode((FONT_DIR / fname).read_bytes()).decode("ascii") for key, fname in names.items()}


def render_html(resume: dict[str, Any]) -> str:
    env = Environment(loader=FileSystemLoader(str(TEMPLATE_DIR)), autoescape=True)
    template = env.get_template(TEMPLATE_NAME)
    return template.render(**resume, **_font_b64())


def render_pdf(resume: dict[str, Any], out_path: Path) -> Path:
    html = render_html(resume)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content(html, wait_until="load")
        page.pdf(path=str(out_path), format="Letter", print_background=True,
                margin={"top": "0in", "bottom": "0in", "left": "0in", "right": "0in"})
        browser.close()
    return out_path


def resume_filename(first: str, last: str, company: str | None = None, ext: str = "pdf") -> str:
    base = f"{first}_{last}_Resume"
    if company:
        safe_company = "".join(c for c in company if c.isalnum() or c in " -_").strip().replace(" ", "")
        base += f"_{safe_company}"
    return f"{base}.{ext}"


_ACCENT = RGBColor(0x1A, 0x1A, 0x1A)
_MUTED = RGBColor(0x44, 0x44, 0x44)


def _section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    pPr = p._p.get_or_add_pPr()
    border = pPr.makeelement("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr")
    bottom = pPr.makeelement("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bottom")
    for attr, val in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1A1A1A")):
        bottom.set(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}", val)
    border.append(bottom)
    pPr.append(border)


def render_docx(resume: dict[str, Any], out_path: Path) -> Path:
    """Mirrors resume.html.j2's layout as closely as Word's model allows: two-column header
    (name+links left, location+contact right), bold job titles on their own line, bulleted
    education/skills - matching Devanshu's actual current resume's format, not a generic one."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Pt(29)
    section.left_margin = section.right_margin = Pt(40)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.font.color.rgb = _ACCENT

    contact = resume["contact"]
    header = doc.add_table(rows=2, cols=2)
    header.autofit = True
    header.columns[0].width = Pt(320)
    header.columns[1].width = Pt(148)

    name_cell = header.cell(0, 0)
    name_p = name_cell.paragraphs[0]
    name_run = name_p.add_run(contact["name"].upper())
    name_run.bold = True
    name_run.font.size = Pt(17)

    links_p = name_cell.add_paragraph()
    links_p.paragraph_format.space_before = Pt(2)
    links_text = " | ".join(b for b in (contact.get("linkedin"), contact.get("portfolio")) if b)
    links_run = links_p.add_run(links_text)
    links_run.font.size = Pt(9)

    loc_cell = header.cell(0, 1)
    loc_p = loc_cell.paragraphs[0]
    loc_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    loc_run = loc_p.add_run(contact.get("location", ""))
    loc_run.bold = True
    loc_run.font.size = Pt(10)

    contact2_p = loc_cell.add_paragraph()
    contact2_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    contact2_p.paragraph_format.space_before = Pt(2)
    contact2_text = " | ".join(b for b in (contact.get("email"), contact.get("phone")) if b)
    contact2_run = contact2_p.add_run(contact2_text)
    contact2_run.font.size = Pt(9)

    if resume.get("summary"):
        _section_heading(doc, "Summary")
        p = doc.add_paragraph(resume["summary"])
        p.paragraph_format.space_after = Pt(0)

    _section_heading(doc, "Experience")
    for e in resume.get("experience", []):
        org_p = doc.add_paragraph()
        org_p.paragraph_format.space_before = Pt(5)
        org_p.paragraph_format.space_after = Pt(0)
        org_p.paragraph_format.tab_stops.add_tab_stop(Pt(468), alignment=WD_TAB_ALIGNMENT.RIGHT)
        org_run = org_p.add_run(e["company"])
        org_run.bold = True
        if e.get("location"):
            org_p.add_run(f" - {e['location']}")
        dates_run = org_p.add_run(f"\t{e.get('start_date', '')} - {e.get('end_date', '')}")
        dates_run.font.size = Pt(9)
        dates_run.font.color.rgb = _MUTED

        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_after = Pt(1)
        title_run = title_p.add_run(e.get("title", ""))
        title_run.bold = True
        title_run.font.size = Pt(9.6)
        for b in e.get("bullets", []):
            bp = doc.add_paragraph(b, style="List Bullet")
            bp.paragraph_format.space_after = Pt(1)

    _section_heading(doc, "Education")
    for ed in resume.get("education", []):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
        head_run = p.add_run(ed["school"])
        head_run.bold = True
        if ed.get("dates"):
            dates_run = p.add_run(f" ({ed['dates']})")
            dates_run.font.size = Pt(8.8)
        p.add_run(" - ")
        deg_run = p.add_run(ed.get("degree", ""))
        if ed.get("coursework"):
            cw_p = doc.add_paragraph()
            cw_p.paragraph_format.left_indent = Pt(18)
            cw_p.paragraph_format.space_after = Pt(4)
            cw_run = cw_p.add_run("Relevant Coursework: " + ", ".join(ed["coursework"]))
            cw_run.font.size = Pt(8.6)
            cw_run.font.color.rgb = _MUTED

    _section_heading(doc, "Skills")
    for cat, items in resume.get("skills", {}).items():
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        cat_run = p.add_run(cat)
        cat_run.bold = True
        p.add_run(" - " + ", ".join(items))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
